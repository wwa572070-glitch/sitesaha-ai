from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Iterable

from .db import KnowledgeBase
from .text import chunk_markdown, chunk_plain_text, normalize_text


SUPPORTED_SUFFIXES = {".md", ".markdown", ".txt", ".pdf", ".docx", ".zip"}


def _markdown_title(text: str) -> str:
    match = re.search(r"^\s{0,3}#\s+(.+?)\s*$", text, flags=re.MULTILINE)
    return match.group(1).strip() if match else ""


def _pdf_text(path: Path) -> list[tuple[str, str]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("读取 PDF 需要安装 pypdf：pip install pypdf") from exc

    reader = PdfReader(str(path))
    pages: list[tuple[str, str]] = []
    for index, page in enumerate(reader.pages, start=1):
        text = normalize_text(page.extract_text() or "")
        if text:
            pages.append((f"第 {index} 页", text))
    if not pages:
        raise ValueError("PDF 未提取到文字，可能是扫描件，请先进行 OCR。")
    return pages


def _docx_text(path: Path) -> tuple[str, str]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("读取 DOCX 需要安装 python-docx：pip install python-docx") from exc

    document = Document(str(path))
    title = ""
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)
        if not text:
            continue
        if not title and paragraph.style and paragraph.style.name.startswith("Heading"):
            title = text
        lines.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        lines.append(f"表格 {table_index}")
        for row in table.rows:
            cells = [normalize_text(cell.text) for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return title, "\n\n".join(lines)


def _zip_member_text(path: Path) -> tuple[str, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            names = [
                name
                for name in archive.namelist()
                if Path(name).suffix.lower() in {".md", ".markdown", ".txt"}
                and not name.startswith("__MACOSX/")
            ]
            if not names:
                raise ValueError("ZIP 中没有可直接读取的 Markdown 或 TXT 文件。")
            parts: list[str] = []
            title = ""
            for name in names:
                text = archive.read(name).decode("utf-8-sig", "replace")
                if not title:
                    title = _markdown_title(text) or Path(name).stem
                parts.append(f"# {Path(name).stem}\n\n{text}")
            return title, "\n\n".join(parts)
    except zipfile.BadZipFile as exc:
        raise ValueError("ZIP 文件格式无效。") from exc


def ingest_file(
    kb: KnowledgeBase,
    path: str | Path,
    *,
    title: str | None = None,
    version: str = "",
    category: str = "",
    source_root: str | Path | None = None,
    replace: bool = True,
) -> tuple[int, int]:
    file_path = Path(path).resolve()
    if not file_path.exists():
        raise FileNotFoundError(file_path)
    if file_path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持的文件类型：{file_path.suffix}")

    if source_root:
        try:
            source = str(file_path.relative_to(Path(source_root).resolve()))
        except ValueError:
            source = file_path.name
    else:
        source = file_path.name

    text = ""
    inferred_title = ""
    if file_path.suffix.lower() == ".docx":
        inferred_title, text = _docx_text(file_path)
    elif file_path.suffix.lower() == ".zip":
        inferred_title, text = _zip_member_text(file_path)
    elif file_path.suffix.lower() != ".pdf":
        text = file_path.read_text(encoding="utf-8-sig")

    if file_path.suffix.lower() in {".md", ".markdown"} and not inferred_title:
        inferred_title = _markdown_title(text)
    document_title = title or inferred_title or file_path.stem
    if file_path.suffix.lower() == ".pdf":
        pages = _pdf_text(file_path)
        combined = normalize_text("\n\n".join(text for _, text in pages))
        chunks = chunk_plain_text(combined, max_chars=850, overlap_chars=80)
    elif file_path.suffix.lower() in {".docx", ".zip"}:
        chunks = chunk_plain_text(text, max_chars=900, overlap_chars=100)
    else:
        if file_path.suffix.lower() in {".md", ".markdown"}:
            chunks = chunk_markdown(text, max_chars=900, overlap_chars=100)
        else:
            chunks = chunk_plain_text(text, max_chars=900, overlap_chars=100)

    return kb.add_document(
        title=document_title,
        source=source,
        version=version,
        category=category,
        chunks=chunks,
        replace=replace,
    )


def ingest_directory(
    kb: KnowledgeBase,
    directory: str | Path,
    *,
    category: str = "",
    replace: bool = True,
) -> list[tuple[Path, int, int]]:
    root = Path(directory).resolve()
    results: list[tuple[Path, int, int]] = []
    if not root.exists():
        return results

    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        document_id, chunk_count = ingest_file(
            kb,
            path,
            category=category,
            source_root=root,
            replace=replace,
        )
        results.append((path, document_id, chunk_count))
    return results


def seed_default_knowledge(
    kb: KnowledgeBase,
    knowledge_dir: str | Path,
    *,
    normalized_sources_dir: str | Path | None = None,
    force: bool = False,
) -> list[tuple[Path, int, int]]:
    if kb.stats()["documents"] and not force:
        return []
    results = ingest_directory(
        kb, knowledge_dir, category="演示摘要", replace=force
    )
    if normalized_sources_dir:
        results.extend(
            ingest_directory(
                kb,
                normalized_sources_dir,
                category="公开标准",
                replace=force,
            )
        )
    return results


def first_existing(paths: Iterable[Path]) -> Path | None:
    return next((path for path in paths if path.exists()), None)
